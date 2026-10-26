import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
ACCENT = RGBColor(0x1F, 0x5F, 0x8B)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.6)
    s.left_margin = s.right_margin = Cm(2.2)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.1


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "1F5F8B")):
        e.set(qn("w:" + k), v)
    b.append(e)
    pPr.append(b)


def line(text, bold=False, after=0, color=None, size=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


# Header (matches CV)
line("Robert Heideman", bold=True, color=ACCENT, size=22, align=WD_ALIGN_PARAGRAPH.CENTER)
line("Bicycle Technician  |  Workshop Lead  |  Skills Mentor", color=GREY, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER)
p = line("Wishart QLD 4122  •  0402 496 428  •  robert@heideman.com.au", size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
bottom_border(p)

line("26 October 2026", after=10)
line("Hiring Team – PUSH! Program")
line("Younity Community Services Ltd")
line("Deception Bay QLD", after=12)
line("Re: Bicycle Technician – PUSH! Program, Deception Bay", bold=True, after=10)

body = [
    "Dear PUSH! Hiring Team,",

    "I'm excited to apply for the Bicycle Technician role with PUSH! Bikes have been a big part of my life for as long "
    "as I can remember, and so has sharing what I know about them. A program where young people build their own bike, "
    "learn real skills and ride away proud of what they've made is exactly the kind of work I'd love to be part of.",

    "I've worked as a bike mechanic for more than 15 years, starting at The Ride Cycles in Kew in 2009. In 2020 I founded Heiden Cycle Works, where I design and build custom "
    "frames from carbon tubes and 3D-printed lugs, and build and maintain all kinds of bikes, from budget commuters to "
    "high-end race bikes. Since 2016 I've also been a travelling tour mechanic with Wide Open Road Cycling, keeping "
    "every bike on the tour safe and running, often from a mobile kit at the side of the road. Inspecting a bike, "
    "working out what it needs and getting it safely back on the road is what I do every day, whether it's a "
    "high-end machine or a well-loved donated bike. I'm comfortable running a workshop on my own and just as happy "
    "working out of a van at a skate park.",

    "My other career, ten years as a Prosthetist-Orthotist and technician, has made me a better teacher. As Senior "
    "Orthotic Technician at Queensland Orthotic Lab, I mentor and train junior and new technicians in a hands-on "
    "workshop, keep tools and equipment in good order, and break complex jobs into simple steps people can master. "
    "Clinical work has also taught me to explain technical things in plain, encouraging language to people of all "
    "ages and backgrounds, and to be patient while someone builds confidence in a new skill.",

    "I'm not a youth worker, but I really enjoy being around young people. For eight years I was the volunteer mechanic "
    "for Warrandyte High School's team on the Great Victorian Bike Ride, working alongside students and teachers to "
    "keep everyone riding. Whether it's showing someone how to tighten their first brake cable or helping them build "
    "a whole bike, I'd love to help young people feel capable and included. I'd also enjoy working with your youth "
    "development team to turn bike-building into fun, hands-on learning sessions.",

    "I hold a current Queensland Blue Card (I've had a Working with Children Check for over 10 years), a current NDIS "
    "Worker Screening Check (Yellow Card) and a current Australian driver licence. I can commit three days a week to "
    "PUSH! and could start from 22 November.",

    "Thank you for considering my application. I'd welcome the chance to chat about how I could contribute to PUSH!, "
    "and I'd be happy to bring along one of my frames to show you what I build.",
]
for t in body:
    p = doc.add_paragraph(t)
    p.paragraph_format.space_after = Pt(8)

line("Kind regards,", after=18)
line("Robert Heideman", bold=True)
line("0402 496 428  •  robert@heideman.com.au", size=10, color=GREY)

doc.save(OUT)
print("saved", OUT)
